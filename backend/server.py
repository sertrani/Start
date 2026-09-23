from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import (FastAPI, APIRouter, HTTPException, Depends, Request, Response,
                     UploadFile, File, Form, Header, Query, BackgroundTasks)
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional, Dict
from datetime import datetime, timezone, timedelta, date
import calendar
import uuid
import io
import re
import ipaddress
import secrets as pysecrets
import logging
import bcrypt
import jwt
import httpx
import requests
from html import escape
from html.parser import HTMLParser
from urllib.parse import urlparse
from bson import ObjectId

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

JWT_ALGORITHM = "HS256"
MAX_SUSPENSION_DAYS = 304
GRACE_DAYS = 15
POLICY_TYPES = ["annuale", "quadrimestrale", "trimestrale", "mensile", "a_data_fissa"]
FRAZIONAMENTO_TYPES = ["unica", "semestrale", "quadrimestrale", "trimestrale", "mensile"]
GRACE_AUTO_TYPES = {"annuale"}
VEHICLE_TYPES = ["auto", "furgone", "altro"]
DOC_TYPES = ["libretto", "carta_circolazione", "polizza", "altro"]

PERMISSIONS = ["manage_vehicles", "manage_policies", "manage_payments",
               "delete_operations", "manage_users", "export_reports", "manage_settings"]
PERMISSION_LABELS = {
    "manage_vehicles": "Gestione veicoli",
    "manage_policies": "Gestione polizze e sospensioni",
    "manage_payments": "Registrazione pagamenti (bollo/collaudo)",
    "delete_operations": "Annullare / eliminare operazioni",
    "manage_users": "Gestione utenti",
    "export_reports": "Esportazione report",
    "manage_settings": "Impostazioni",
}

def today_iso() -> str:
    return datetime.now(timezone.utc).date().isoformat()

def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()

# ---------------- Storage ----------------
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_NAME = "fleetcare"
storage_key = None

def init_storage(force: bool = False):
    global storage_key
    if storage_key and not force:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key

def put_object(path: str, data: bytes, content_type: str) -> dict:
    last = None
    for attempt in range(3):
        key = init_storage(force=(attempt > 0))
        resp = requests.put(f"{STORAGE_URL}/objects/{path}",
                            headers={"X-Storage-Key": key, "Content-Type": content_type},
                            data=data, timeout=120)
        if resp.status_code < 400:
            return resp.json()
        last = resp
        if resp.status_code not in (404, 500, 503):
            break
    last.raise_for_status()

def get_object(path: str):
    last = None
    for attempt in range(3):
        key = init_storage(force=(attempt > 0))
        resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
        if resp.status_code < 400:
            return resp.content, resp.headers.get("Content-Type", "application/octet-stream")
        last = resp
        if resp.status_code not in (404, 500, 503):
            break
    last.raise_for_status()

MIME_TYPES = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
              "gif": "image/gif", "webp": "image/webp", "pdf": "application/pdf"}
ALLOWED_DOC_EXT = {"pdf", "jpg", "jpeg", "png", "webp"}

def is_valid_image_bytes(data: bytes) -> bool:
    try:
        from PIL import Image as PILImage
        im = PILImage.open(io.BytesIO(data))
        im.load()
        return True
    except Exception:
        return False

# ---------------- Email ----------------
EMAIL_BASE_URL = "https://integrations.emergentagent.com"
EMAIL_KEY = os.environ.get("EMERGENT_EMAIL_KEY")
EMAIL_FROM_NAME = os.environ.get("EMAIL_FROM_NAME", "FleetCare Autonoleggio")
EMAIL_REPLY_TO = os.environ.get("EMAIL_REPLY_TO")
_SHORTENERS = ("bit.ly", "tinyurl.com", "t.co", "is.gd", "cutt.ly", "goo.gl", "rebrand.ly")
_CRED_ASK = ("reply with your password", "send your password", "cvv", "seed phrase",
             "recovery phrase", "confirm your card number", "social security number")
_HOSTISH = re.compile(r"\b(?:https?://)?((?:[a-z0-9-]+\.)+[a-z]{2,})", re.I)

def _host_ok(host: str) -> bool:
    if not host or "xn--" in host:
        return False
    try:
        ipaddress.ip_address(host)
        return False
    except ValueError:
        pass
    return not any(host == s or host.endswith("." + s) for s in _SHORTENERS)

def _same_site(shown: str, real: str) -> bool:
    return shown == real or real.endswith("." + shown) or shown.endswith("." + real)

class _EmailScan(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags, self.urls, self.anchors = set(), [], []
        self._href, self._text = None, []
    def handle_starttag(self, tag, attrs):
        self.tags.add(tag.lower())
        self.urls += [v for k, v in attrs if k.lower() in ("href", "src") and v]
        if tag.lower() == "a":
            self._href = dict((k.lower(), v) for k, v in attrs).get("href")
            self._text = []
    def handle_data(self, data):
        if self._href is not None:
            self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == "a" and self._href is not None:
            self.anchors.append((self._href, "".join(self._text)))
            self._href, self._text = None, []

def _assert_safe_email(subject: str, html: str) -> None:
    scan = _EmailScan(); scan.feed(html)
    if scan.tags & {"form", "input", "textarea", "select"}:
        raise ValueError("No forms in email (G2)")
    body = f"{subject}\n{html}".lower()
    for p in _CRED_ASK:
        if p in body:
            raise ValueError(f"Email asks for credentials: {p!r} (G2)")
    for url in scan.urls:
        low = url.strip().lower()
        if low.startswith(("mailto:", "tel:", "cid:", "#")):
            continue
        if not low.startswith("https://"):
            raise ValueError(f"Email links must be https: {url!r} (G3)")
        host = urlparse(low).hostname or ""
        if not _host_ok(host) or urlparse(low).username is not None:
            raise ValueError(f"Unsafe URL: {url!r} (G3)")
    for href, text in scan.anchors:
        real = urlparse(href.strip().lower()).hostname or ""
        if not real:
            continue
        for m in _HOSTISH.finditer(text):
            if not _same_site(m.group(1).lower(), real):
                raise ValueError(f"Anchor text mismatch (G3)")

async def send_email(*, to: str, subject: str, html: str) -> Optional[str]:
    _assert_safe_email(subject, html)
    payload = {"to": [to], "subject": subject, "html": html, "from_name": EMAIL_FROM_NAME}
    if EMAIL_REPLY_TO:
        payload["contact_email"] = EMAIL_REPLY_TO
    async with httpx.AsyncClient(timeout=30) as c:
        resp = await c.post(f"{EMAIL_BASE_URL}/api/v1/email/send",
                            headers={"X-Email-Key": EMAIL_KEY}, json=payload)
    resp.raise_for_status()
    return resp.json().get("id")

# ---------------- Auth ----------------
def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")

def verify_password(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))

def get_jwt_secret() -> str:
    return os.environ["JWT_SECRET"]

def create_access_token(user_id: str, email: str) -> str:
    payload = {"sub": user_id, "email": email,
               "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "access"}
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)

def user_perms(user: dict) -> set:
    if user.get("role") == "admin":
        return set(PERMISSIONS)
    return set(user.get("permissions", []))

def public_user(u: dict) -> dict:
    return {"id": str(u["_id"]) if not isinstance(u.get("_id"), str) else u["_id"],
            "email": u["email"], "name": u.get("name"), "role": u.get("role", "user"),
            "permissions": sorted(user_perms(u)), "is_active": u.get("is_active", True)}

async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth_header = request.headers.get("Authorization", "")
        if auth_header.startswith("Bearer "):
            token = auth_header[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Non autenticato")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "access":
            raise HTTPException(status_code=401, detail="Token non valido")
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(status_code=401, detail="Utente non trovato")
        if not user.get("is_active", True):
            raise HTTPException(status_code=403, detail="Account disattivato")
        user["_id"] = str(user["_id"])
        user.pop("password_hash", None)
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token scaduto")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token non valido")

def require(*perms):
    async def dep(user: dict = Depends(get_current_user)):
        up = user_perms(user)
        if not all(p in up for p in perms):
            raise HTTPException(status_code=403, detail="Permesso negato per questa operazione")
        return user
    return dep

# ---------------- Models ----------------
class LoginInput(BaseModel):
    email: EmailStr
    password: str

class Policy(BaseModel):
    compagnia: str
    tipologia: str
    numero_polizza: Optional[str] = None
    data_stipula: str
    scadenza_rata_intermedia: Optional[str] = None
    scadenza_contratto: str
    importo_premio: Optional[float] = None
    importo_rata: Optional[float] = None
    grace_period: Optional[bool] = False
    frazionamento: Optional[str] = None

class RenewInput(Policy):
    reset_suspensions: bool = False

class VehicleInput(BaseModel):
    targa: str
    marca_modello: str
    data_immatricolazione: str
    bollo_scadenza: Optional[str] = None
    note: Optional[str] = Field(None, max_length=2000)
    tipo: Optional[str] = "auto"

class CollaudoInput(BaseModel):
    data_collaudo: str
    note: Optional[str] = None

class BolloPaymentInput(BaseModel):
    data_pagamento: str
    importo: float
    periodo: Optional[str] = None
    nuova_scadenza: Optional[str] = None
    note: Optional[str] = None

class SuspendInput(BaseModel):
    effective_date: Optional[str] = None

class NotificationDays(BaseModel):
    bollo: int = 30
    collaudo: int = 30
    polizza: int = 30

class SettingsInput(BaseModel):
    company_name: Optional[str] = None
    notification_recipients: List[str] = []
    notification_days: NotificationDays = NotificationDays()
    bell_days: int = Field(7, ge=0, le=365)
    season_start_month: int = Field(4, ge=1, le=12)
    season_end_month: int = Field(10, ge=1, le=12)

class UserCreate(BaseModel):
    email: EmailStr
    password: str
    name: str
    permissions: List[str] = []

class UserUpdate(BaseModel):
    name: Optional[str] = None
    permissions: Optional[List[str]] = None
    is_active: Optional[bool] = None

class PasswordResetInput(BaseModel):
    password: str

# ---------------- Date / business logic ----------------
def parse_date(s: Optional[str]) -> Optional[date]:
    if not s:
        return None
    return date.fromisoformat(s[:10])

def end_of_month(year: int, month: int) -> date:
    return date(year, month, calendar.monthrange(year, month)[1])

def compute_collaudo_deadline(immatricolazione: str, last_collaudo: Optional[str]) -> Optional[str]:
    lc = parse_date(last_collaudo)
    if lc:
        return end_of_month(lc.year + 2, lc.month).isoformat()
    im = parse_date(immatricolazione)
    if im:
        return end_of_month(im.year + 4, im.month).isoformat()
    return None

def days_status(deadline: Optional[str]):
    if not deadline:
        return ("none", None)
    d = parse_date(deadline)
    today = datetime.now(timezone.utc).date()
    delta = (d - today).days
    if delta < 0:
        return ("expired", delta)
    if delta <= 30:
        return ("upcoming", delta)
    return ("valid", delta)

def normalize_policy_suspension(policy: dict) -> bool:
    changed = False
    cumulative = policy.get("cumulative_suspension_days", 0)
    if policy.get("status") == "suspended" and policy.get("current_suspension_start"):
        start = parse_date(policy["current_suspension_start"])
        today = datetime.now(timezone.utc).date()
        ongoing = (today - start).days
        if cumulative + ongoing >= MAX_SUSPENSION_DAYS:
            used = MAX_SUSPENSION_DAYS - cumulative
            policy.setdefault("suspensions", []).append({
                "id": str(uuid.uuid4()), "suspended_at": policy["current_suspension_start"],
                "reactivated_at": today.isoformat(), "days": used, "auto": True})
            policy["cumulative_suspension_days"] = MAX_SUSPENSION_DAYS
            policy["status"] = "active"
            policy["current_suspension_start"] = None
            policy["suspension_limit_reached"] = True
            changed = True
    return changed

def effective_suspension_days(policy: dict) -> int:
    cumulative = policy.get("cumulative_suspension_days", 0)
    if policy.get("status") == "suspended" and policy.get("current_suspension_start"):
        start = parse_date(policy["current_suspension_start"])
        today = datetime.now(timezone.utc).date()
        return min(cumulative + (today - start).days, MAX_SUSPENSION_DAYS)
    return min(cumulative, MAX_SUSPENSION_DAYS)

def grace_applies_to(policy: dict) -> bool:
    return policy.get("tipologia") in GRACE_AUTO_TYPES or bool(policy.get("grace_period"))

def fresh_suspension_state() -> dict:
    return {"status": "active", "cumulative_suspension_days": 0, "current_suspension_start": None,
            "suspensions": [], "suspension_limit_reached": False}

def compute_last_collaudo(history: list) -> Optional[str]:
    dates = [h["data_collaudo"] for h in history if h.get("data_collaudo")]
    return max(dates) if dates else None

def build_vehicle_view(v: dict) -> dict:
    v.pop("_id", None)
    history = v.get("collaudo_history", [])
    last_collaudo = compute_last_collaudo(history) or v.get("last_collaudo_date")
    collaudo_deadline = compute_collaudo_deadline(v.get("data_immatricolazione"), last_collaudo)
    collaudo_state, collaudo_days = days_status(collaudo_deadline)
    bollo_state, bollo_days = days_status(v.get("bollo_scadenza"))

    policy = v.get("policy")
    insurance_state = "none"
    insurance_days = None
    policy_view = None
    insurance_ok = False
    today = datetime.now(timezone.utc).date()
    if policy:
        contract = parse_date(policy.get("scadenza_contratto"))
        rata_state, rata_days = days_status(policy.get("scadenza_rata_intermedia"))
        eff_days = effective_suspension_days(policy)
        suspended = policy.get("status") == "suspended"
        limit_reached = policy.get("suspension_limit_reached", False)
        g_applies = grace_applies_to(policy)
        grace_end = contract + timedelta(days=GRACE_DAYS) if (contract and g_applies) else contract
        if suspended:
            insurance_state = "suspended"
            insurance_ok = False
        else:
            base_state, base_days = days_status(policy.get("scadenza_contratto"))
            if contract and today > contract and grace_end and today <= grace_end:
                insurance_state = "grace"; insurance_ok = True; insurance_days = (grace_end - today).days
            elif grace_end and today > grace_end:
                insurance_state = "expired"; insurance_ok = False
                insurance_days = (contract - today).days if contract else None
            else:
                insurance_state = base_state; insurance_ok = base_state != "expired"; insurance_days = base_days
        policy_view = {
            "compagnia": policy.get("compagnia"), "tipologia": policy.get("tipologia"),
            "numero_polizza": policy.get("numero_polizza"),
            "data_stipula": policy.get("data_stipula"),
            "scadenza_rata_intermedia": policy.get("scadenza_rata_intermedia"),
            "scadenza_contratto": policy.get("scadenza_contratto"),
            "importo_premio": policy.get("importo_premio"), "importo_rata": policy.get("importo_rata"),
            "frazionamento": policy.get("frazionamento"),
            "grace_period": bool(policy.get("grace_period")), "grace_applies": g_applies,
            "grace_end": grace_end.isoformat() if grace_end else None,
            "status": policy.get("status", "active"), "rata_state": rata_state, "rata_days": rata_days,
            "cumulative_suspension_days": eff_days, "max_suspension_days": MAX_SUSPENSION_DAYS,
            "suspension_limit_reached": limit_reached,
            "current_suspension_start": policy.get("current_suspension_start"),
            "suspensions": policy.get("suspensions", []),
            "can_suspend": (not suspended) and (not limit_reached) and eff_days < MAX_SUSPENSION_DAYS,
        }

    collaudo_ok = collaudo_state != "expired"
    can_circulate = insurance_ok and collaudo_ok
    docs = [d for d in v.get("documents", []) if not d.get("is_deleted")]
    return {
        **{k: val for k, val in v.items() if k not in ("documents",)},
        "last_collaudo_date": last_collaudo,
        "collaudo_deadline": collaudo_deadline, "collaudo_state": collaudo_state, "collaudo_days": collaudo_days,
        "collaudo_history": history,
        "bollo_state": bollo_state, "bollo_days": bollo_days, "bollo_history": v.get("bollo_history", []),
        "insurance_state": insurance_state, "insurance_days": insurance_days,
        "policy": policy_view, "policy_archive": v.get("policy_archive", []),
        "documents": docs, "can_circulate": can_circulate,
        "collaudo_ok": collaudo_ok, "insurance_ok": insurance_ok,
    }

async def get_vehicle_or_404(vehicle_id: str) -> dict:
    v = await db.vehicles.find_one({"id": vehicle_id})
    if not v:
        raise HTTPException(status_code=404, detail="Veicolo non trovato")
    if v.get("policy") and normalize_policy_suspension(v["policy"]):
        await db.vehicles.update_one({"id": vehicle_id}, {"$set": {"policy": v["policy"]}})
    return v

async def get_settings() -> dict:
    s = await db.settings.find_one({"key": "app"})
    if not s:
        s = {"key": "app", "company_name": "FleetCare Autonoleggio", "notification_recipients": [],
             "notification_days": {"bollo": 30, "collaudo": 30, "polizza": 30}, "bell_days": 7,
             "season_start_month": 4, "season_end_month": 10, "logo_path": None}
        await db.settings.insert_one(dict(s))
    s.pop("_id", None)
    if "notification_days" not in s:
        legacy = s.get("notification_days_before", 30)
        s["notification_days"] = {"bollo": legacy, "collaudo": legacy, "polizza": legacy}
    if "bell_days" not in s:
        s["bell_days"] = 7
    s.setdefault("season_start_month", 4)
    s.setdefault("season_end_month", 10)
    return s

# ---------------- Audit ----------------
async def log_op(user: dict, action: str, label: str, vehicle: Optional[dict],
                 description: str, restore: Optional[dict], effective_date: Optional[str] = None):
    entry = {
        "id": str(uuid.uuid4()), "ts": now_iso(),
        "effective_date": effective_date or today_iso(),
        "user_id": user["_id"], "user_email": user["email"], "user_name": user.get("name"),
        "action": action, "action_label": label,
        "vehicle_id": (vehicle or {}).get("id"), "targa": (vehicle or {}).get("targa"),
        "marca_modello": (vehicle or {}).get("marca_modello"),
        "description": description, "restore": restore, "reverted": False,
    }
    await db.audit.insert_one(dict(entry))
    entry.pop("_id", None)
    return entry

def snapshot(v: dict, fields: list) -> dict:
    return {f: v.get(f) for f in fields}

# ---------------- Auth endpoints ----------------
@api_router.post("/auth/login")
async def login(input: LoginInput, response: Response, request: Request):
    email = input.email.lower()
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(input.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Credenziali non valide")
    if not user.get("is_active", True):
        raise HTTPException(status_code=403, detail="Account disattivato")
    uid = str(user["_id"])
    token = create_access_token(uid, email)
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=604800, path="/")
    xff = request.headers.get("x-forwarded-for")
    ip = xff.split(",")[0].strip() if xff else (request.client.host if request.client else None)
    await db.login_log.insert_one({"id": str(uuid.uuid4()), "user_id": uid, "email": email,
                                   "name": user.get("name"), "role": user.get("role", "user"),
                                   "ts": now_iso(), "ip": ip})
    return {"token": token, "user": public_user(user)}

@api_router.post("/auth/logout")
async def logout(response: Response, user: dict = Depends(get_current_user)):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}

@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return public_user(user)

@api_router.get("/auth/permissions")
async def list_permissions(user: dict = Depends(get_current_user)):
    return {"permissions": PERMISSIONS, "labels": PERMISSION_LABELS}

# ---------------- Users (admin) ----------------
@api_router.get("/users")
async def list_users(user: dict = Depends(require("manage_users"))):
    users = await db.users.find().sort("created_at", 1).to_list(500)
    return [public_user(u) for u in users]

@api_router.post("/users")
async def create_user(input: UserCreate, user: dict = Depends(require("manage_users"))):
    email = input.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email già registrata")
    perms = [p for p in input.permissions if p in PERMISSIONS]
    doc = {"email": email, "password_hash": hash_password(input.password), "name": input.name,
           "role": "user", "permissions": perms, "is_active": True,
           "created_at": now_iso(), "created_by": user["email"]}
    res = await db.users.insert_one(doc)
    doc["_id"] = res.inserted_id
    return public_user(doc)

@api_router.put("/users/{user_id}")
async def update_user(user_id: str, input: UserUpdate, user: dict = Depends(require("manage_users"))):
    target = await db.users.find_one({"_id": ObjectId(user_id)})
    if not target:
        raise HTTPException(status_code=404, detail="Utente non trovato")
    upd = {}
    if input.name is not None:
        upd["name"] = input.name
    if input.permissions is not None and target.get("role") != "admin":
        upd["permissions"] = [p for p in input.permissions if p in PERMISSIONS]
    if input.is_active is not None:
        if target.get("role") == "admin" and not input.is_active:
            raise HTTPException(status_code=400, detail="Non puoi disattivare l'amministratore")
        upd["is_active"] = input.is_active
    if upd:
        await db.users.update_one({"_id": ObjectId(user_id)}, {"$set": upd})
    return public_user(await db.users.find_one({"_id": ObjectId(user_id)}))

@api_router.post("/users/{user_id}/reset-password")
async def reset_user_password(user_id: str, input: PasswordResetInput, user: dict = Depends(require("manage_users"))):
    target = await db.users.find_one({"_id": ObjectId(user_id)})
    if not target:
        raise HTTPException(status_code=404, detail="Utente non trovato")
    await db.users.update_one({"_id": ObjectId(user_id)}, {"$set": {"password_hash": hash_password(input.password)}})
    return {"ok": True}

@api_router.delete("/users/{user_id}")
async def delete_user(user_id: str, user: dict = Depends(require("manage_users"))):
    target = await db.users.find_one({"_id": ObjectId(user_id)})
    if not target:
        raise HTTPException(status_code=404, detail="Utente non trovato")
    if target.get("role") == "admin":
        raise HTTPException(status_code=400, detail="Non puoi eliminare l'amministratore")
    if str(target["_id"]) == user["_id"]:
        raise HTTPException(status_code=400, detail="Non puoi eliminare te stesso")
    await db.users.delete_one({"_id": ObjectId(user_id)})
    return {"ok": True}

# ---------------- Vehicles ----------------
@api_router.get("/vehicles")
async def list_vehicles(user: dict = Depends(get_current_user)):
    vehicles = await db.vehicles.find().sort("created_at", -1).to_list(2000)
    out = []
    for v in vehicles:
        if v.get("policy") and normalize_policy_suspension(v["policy"]):
            await db.vehicles.update_one({"id": v["id"]}, {"$set": {"policy": v["policy"]}})
        out.append(build_vehicle_view(v))
    return out

@api_router.post("/vehicles")
async def create_vehicle(input: VehicleInput, user: dict = Depends(require("manage_vehicles"))):
    doc = input.model_dump()
    doc["targa"] = doc["targa"].upper().strip()
    doc["id"] = str(uuid.uuid4())
    doc["policy"] = None
    doc["policy_archive"] = []
    doc["documents"] = []
    doc["bollo_history"] = []
    doc["collaudo_history"] = []
    doc["created_at"] = now_iso()
    await db.vehicles.insert_one(doc)
    await log_op(user, "vehicle_create", "Creazione veicolo", doc,
                 f"Aggiunto veicolo {doc['targa']} — {doc['marca_modello']}",
                 {"type": "delete_vehicle"})
    return build_vehicle_view(await db.vehicles.find_one({"id": doc["id"]}))

@api_router.get("/vehicles/{vehicle_id}")
async def get_vehicle(vehicle_id: str, user: dict = Depends(get_current_user)):
    return build_vehicle_view(await get_vehicle_or_404(vehicle_id))

@api_router.put("/vehicles/{vehicle_id}")
async def update_vehicle(vehicle_id: str, input: VehicleInput, user: dict = Depends(require("manage_vehicles"))):
    v = await get_vehicle_or_404(vehicle_id)
    prev = snapshot(v, ["targa", "marca_modello", "data_immatricolazione", "bollo_scadenza", "note", "tipo"])
    upd = input.model_dump()
    upd["targa"] = upd["targa"].upper().strip()
    await db.vehicles.update_one({"id": vehicle_id}, {"$set": upd})
    await log_op(user, "vehicle_update", "Modifica veicolo", {**v, **upd},
                 f"Modificati dati veicolo {upd['targa']}", {"type": "set", "set": prev})
    return build_vehicle_view(await db.vehicles.find_one({"id": vehicle_id}))

@api_router.delete("/vehicles/{vehicle_id}")
async def delete_vehicle(vehicle_id: str, user: dict = Depends(require("delete_operations"))):
    v = await get_vehicle_or_404(vehicle_id)
    full = await db.vehicles.find_one({"id": vehicle_id})
    full.pop("_id", None)
    await db.vehicles.delete_one({"id": vehicle_id})
    await log_op(user, "vehicle_delete", "Eliminazione veicolo", v,
                 f"Eliminato veicolo {v.get('targa')}", {"type": "insert_vehicle", "doc": full})
    return {"ok": True}

# ---------------- Bollo ----------------
@api_router.post("/vehicles/{vehicle_id}/bollo")
async def add_bollo_payment(vehicle_id: str, input: BolloPaymentInput, user: dict = Depends(require("manage_payments"))):
    v = await get_vehicle_or_404(vehicle_id)
    prev = snapshot(v, ["bollo_history", "bollo_scadenza"])
    payment = {"id": str(uuid.uuid4()), "data_pagamento": input.data_pagamento, "importo": input.importo,
               "periodo": input.periodo, "note": input.note, "created_by": user["email"], "created_at": now_iso()}
    upd = {"$push": {"bollo_history": payment}}
    if input.nuova_scadenza:
        upd["$set"] = {"bollo_scadenza": input.nuova_scadenza}
    await db.vehicles.update_one({"id": vehicle_id}, upd)
    await log_op(user, "bollo_add", "Pagamento bollo", v,
                 f"Registrato pagamento bollo di € {input.importo:.2f}" + (f" · nuova scadenza {input.nuova_scadenza}" if input.nuova_scadenza else ""),
                 {"type": "set", "set": {"bollo_history": prev["bollo_history"] or [], "bollo_scadenza": prev["bollo_scadenza"]}},
                 effective_date=input.data_pagamento)
    return build_vehicle_view(await db.vehicles.find_one({"id": vehicle_id}))

@api_router.delete("/vehicles/{vehicle_id}/bollo/{payment_id}")
async def delete_bollo_payment(vehicle_id: str, payment_id: str, user: dict = Depends(require("delete_operations"))):
    v = await get_vehicle_or_404(vehicle_id)
    prev = snapshot(v, ["bollo_history"])
    await db.vehicles.update_one({"id": vehicle_id}, {"$pull": {"bollo_history": {"id": payment_id}}})
    await log_op(user, "bollo_delete", "Eliminazione pagamento bollo", v,
                 "Eliminato un pagamento bollo", {"type": "set", "set": {"bollo_history": prev["bollo_history"] or []}})
    return build_vehicle_view(await db.vehicles.find_one({"id": vehicle_id}))

# ---------------- Collaudo ----------------
@api_router.post("/vehicles/{vehicle_id}/collaudo")
async def add_collaudo(vehicle_id: str, input: CollaudoInput, user: dict = Depends(require("manage_payments"))):
    v = await get_vehicle_or_404(vehicle_id)
    prev = snapshot(v, ["collaudo_history", "last_collaudo_date"])
    record = {"id": str(uuid.uuid4()), "data_collaudo": input.data_collaudo, "note": input.note,
              "created_by": user["email"], "created_at": now_iso()}
    history = (v.get("collaudo_history") or []) + [record]
    last = compute_last_collaudo(history)
    await db.vehicles.update_one({"id": vehicle_id}, {"$set": {"collaudo_history": history, "last_collaudo_date": last}})
    new_deadline = compute_collaudo_deadline(v.get("data_immatricolazione"), last)
    await log_op(user, "collaudo_add", "Collaudo eseguito", v,
                 f"Registrato collaudo del {input.data_collaudo} · prossima scadenza {new_deadline}",
                 {"type": "set", "set": {"collaudo_history": prev["collaudo_history"] or [], "last_collaudo_date": prev["last_collaudo_date"]}},
                 effective_date=input.data_collaudo)
    return build_vehicle_view(await db.vehicles.find_one({"id": vehicle_id}))

@api_router.delete("/vehicles/{vehicle_id}/collaudo/{record_id}")
async def delete_collaudo(vehicle_id: str, record_id: str, user: dict = Depends(require("delete_operations"))):
    v = await get_vehicle_or_404(vehicle_id)
    prev = snapshot(v, ["collaudo_history", "last_collaudo_date"])
    history = [h for h in (v.get("collaudo_history") or []) if h.get("id") != record_id]
    last = compute_last_collaudo(history)
    await db.vehicles.update_one({"id": vehicle_id}, {"$set": {"collaudo_history": history, "last_collaudo_date": last}})
    await log_op(user, "collaudo_delete", "Eliminazione collaudo", v, "Eliminato un collaudo",
                 {"type": "set", "set": {"collaudo_history": prev["collaudo_history"] or [], "last_collaudo_date": prev["last_collaudo_date"]}})
    return build_vehicle_view(await db.vehicles.find_one({"id": vehicle_id}))

# ---------------- Policy ----------------
def build_policy(input: Policy, base: dict) -> dict:
    return {"compagnia": input.compagnia, "tipologia": input.tipologia,
            "numero_polizza": input.numero_polizza, "data_stipula": input.data_stipula,
            "scadenza_rata_intermedia": input.scadenza_rata_intermedia,
            "scadenza_contratto": input.scadenza_contratto, "importo_premio": input.importo_premio,
            "importo_rata": input.importo_rata, "grace_period": bool(input.grace_period),
            "frazionamento": (input.frazionamento if input.tipologia == "annuale" else None),
            "status": base.get("status", "active"),
            "cumulative_suspension_days": base.get("cumulative_suspension_days", 0),
            "current_suspension_start": base.get("current_suspension_start"),
            "suspensions": base.get("suspensions", []),
            "suspension_limit_reached": base.get("suspension_limit_reached", False)}

@api_router.put("/vehicles/{vehicle_id}/policy")
async def set_policy(vehicle_id: str, input: Policy, user: dict = Depends(require("manage_policies"))):
    if input.tipologia not in POLICY_TYPES:
        raise HTTPException(status_code=400, detail="Tipologia polizza non valida")
    v = await get_vehicle_or_404(vehicle_id)
    prev = snapshot(v, ["policy", "policy_archive"])
    existing = v.get("policy") or {}
    policy = build_policy(input, existing)
    await db.vehicles.update_one({"id": vehicle_id}, {"$set": {"policy": policy}})
    await log_op(user, "policy_set", "Polizza salvata", v,
                 f"Polizza {input.compagnia}" + (f" n. {input.numero_polizza}" if input.numero_polizza else "") + f" · scad. {input.scadenza_contratto}",
                 {"type": "set", "set": {"policy": prev["policy"], "policy_archive": prev["policy_archive"] or []}},
                 effective_date=input.data_stipula)
    return build_vehicle_view(await db.vehicles.find_one({"id": vehicle_id}))

@api_router.post("/vehicles/{vehicle_id}/policy/renew")
async def renew_policy(vehicle_id: str, input: RenewInput, user: dict = Depends(require("manage_policies"))):
    if input.tipologia not in POLICY_TYPES:
        raise HTTPException(status_code=400, detail="Tipologia polizza non valida")
    v = await get_vehicle_or_404(vehicle_id)
    prev = snapshot(v, ["policy", "policy_archive"])
    old = v.get("policy")
    archive = list(v.get("policy_archive") or [])
    if old:
        archive.append({**old, "archived_at": now_iso(), "archived_by": user["email"]})
    base = fresh_suspension_state() if (input.reset_suspensions or not old) else old
    policy = build_policy(input, base)
    await db.vehicles.update_one({"id": vehicle_id}, {"$set": {"policy": policy, "policy_archive": archive}})
    msg = f"Nuova polizza {input.compagnia}" + (f" n. {input.numero_polizza}" if input.numero_polizza else "")
    msg += " · contatore sospensioni azzerato" if input.reset_suspensions else " · contatore sospensioni mantenuto"
    await log_op(user, "policy_renew", "Rinnovo / nuova polizza", v, msg,
                 {"type": "set", "set": {"policy": prev["policy"], "policy_archive": prev["policy_archive"] or []}},
                 effective_date=input.data_stipula)
    return build_vehicle_view(await db.vehicles.find_one({"id": vehicle_id}))

@api_router.post("/vehicles/{vehicle_id}/policy/suspend")
async def suspend_policy(vehicle_id: str, input: SuspendInput = SuspendInput(), user: dict = Depends(require("manage_policies"))):
    v = await get_vehicle_or_404(vehicle_id)
    policy = v.get("policy")
    if not policy:
        raise HTTPException(status_code=400, detail="Nessuna polizza da sospendere")
    if policy.get("status") == "suspended":
        raise HTTPException(status_code=400, detail="Polizza già sospesa")
    if policy.get("suspension_limit_reached") or policy.get("cumulative_suspension_days", 0) >= MAX_SUSPENSION_DAYS:
        raise HTTPException(status_code=400, detail="Limite massimo di sospensione (10 mesi) raggiunto")
    prev = snapshot(v, ["policy"])
    eff = input.effective_date or today_iso()
    policy["status"] = "suspended"
    policy["current_suspension_start"] = eff
    await db.vehicles.update_one({"id": vehicle_id}, {"$set": {"policy": policy}})
    await log_op(user, "policy_suspend", "Sospensione copertura", v,
                 f"Copertura sospesa dal {eff}", {"type": "set", "set": {"policy": prev["policy"]}}, effective_date=eff)
    return build_vehicle_view(await db.vehicles.find_one({"id": vehicle_id}))

@api_router.post("/vehicles/{vehicle_id}/policy/reactivate")
async def reactivate_policy(vehicle_id: str, input: SuspendInput = SuspendInput(), user: dict = Depends(require("manage_policies"))):
    v = await get_vehicle_or_404(vehicle_id)
    policy = v.get("policy")
    if not policy or policy.get("status") != "suspended":
        raise HTTPException(status_code=400, detail="La polizza non è sospesa")
    prev = snapshot(v, ["policy"])
    start = parse_date(policy["current_suspension_start"])
    eff = input.effective_date or today_iso()
    end = parse_date(eff)
    if end < start:
        raise HTTPException(status_code=400, detail="La data di riattivazione precede la sospensione")
    days = (end - start).days
    cumulative = policy.get("cumulative_suspension_days", 0)
    new_cumulative = min(cumulative + days, MAX_SUSPENSION_DAYS)
    used = new_cumulative - cumulative
    policy.setdefault("suspensions", []).append({
        "id": str(uuid.uuid4()), "suspended_at": policy["current_suspension_start"],
        "reactivated_at": eff, "days": used, "auto": False})
    policy["cumulative_suspension_days"] = new_cumulative
    policy["status"] = "active"
    policy["current_suspension_start"] = None
    if new_cumulative >= MAX_SUSPENSION_DAYS:
        policy["suspension_limit_reached"] = True
    await db.vehicles.update_one({"id": vehicle_id}, {"$set": {"policy": policy}})
    await log_op(user, "policy_reactivate", "Riattivazione copertura", v,
                 f"Copertura riattivata dal {eff} · +{used} gg (totale {new_cumulative}/{MAX_SUSPENSION_DAYS})",
                 {"type": "set", "set": {"policy": prev["policy"]}}, effective_date=eff)
    return build_vehicle_view(await db.vehicles.find_one({"id": vehicle_id}))

# ---------------- Documents ----------------
@api_router.post("/vehicles/{vehicle_id}/documents")
async def upload_document(vehicle_id: str, doc_type: str = Form(...), file: UploadFile = File(...),
                          user: dict = Depends(require("manage_vehicles"))):
    await get_vehicle_or_404(vehicle_id)
    if doc_type not in DOC_TYPES:
        doc_type = "altro"
    ext = (file.filename.rsplit(".", 1)[-1] if "." in file.filename else "bin").lower()
    if ext not in ALLOWED_DOC_EXT:
        raise HTTPException(status_code=400, detail="Formato non consentito (solo PDF, JPG, PNG, WEBP)")
    data = await file.read()
    path = f"{APP_NAME}/vehicles/{vehicle_id}/{uuid.uuid4()}.{ext}"
    ct = MIME_TYPES.get(ext, file.content_type or "application/octet-stream")
    result = put_object(path, data, ct)
    doc = {"id": str(uuid.uuid4()), "doc_type": doc_type, "original_filename": file.filename,
           "storage_path": result["path"], "content_type": ct, "size": result.get("size", len(data)),
           "is_deleted": False, "created_at": now_iso()}
    await db.vehicles.update_one({"id": vehicle_id}, {"$push": {"documents": doc}})
    return doc

@api_router.delete("/vehicles/{vehicle_id}/documents/{doc_id}")
async def delete_document(vehicle_id: str, doc_id: str, user: dict = Depends(require("manage_vehicles"))):
    await get_vehicle_or_404(vehicle_id)
    await db.vehicles.update_one({"id": vehicle_id, "documents.id": doc_id}, {"$set": {"documents.$.is_deleted": True}})
    return {"ok": True}

async def _authorize(auth_header: Optional[str]):
    if not auth_header or not auth_header.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Non autenticato")
    token = auth_header[7:]
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "access":
            raise HTTPException(status_code=401, detail="Token non valido")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token non valido")

@api_router.get("/files/{path:path}")
async def download_file(path: str, authorization: str = Header(None), auth: str = Query(None)):
    await _authorize(authorization or (f"Bearer {auth}" if auth else None))
    data, content_type = get_object(path)
    return Response(content=data, media_type=content_type)

# ---------------- Audit / history ----------------
@api_router.get("/audit")
async def list_audit(vehicle_id: Optional[str] = Query(None), action: Optional[str] = Query(None),
                     user_email: Optional[str] = Query(None),
                     date_from: Optional[str] = Query(None), date_to: Optional[str] = Query(None),
                     limit: int = Query(300), user: dict = Depends(get_current_user)):
    q = build_audit_query(vehicle_id, action, date_from, date_to, user_email)
    entries = await db.audit.find(q).sort("ts", -1).to_list(min(limit, 2000))
    for e in entries:
        e.pop("_id", None)
        e.pop("restore", None)
    return entries

def build_audit_query(vehicle_id, action, date_from, date_to, user_email=None):
    q = {}
    if vehicle_id:
        q["vehicle_id"] = vehicle_id
    if action:
        q["action"] = action
    if user_email:
        q["user_email"] = user_email
    if date_from or date_to:
        rng = {}
        if date_from:
            rng["$gte"] = date_from[:10]
        if date_to:
            rng["$lte"] = date_to[:10] + "T23:59:59.999999+00:00"
        q["ts"] = rng
    return q

@api_router.get("/vehicles/{vehicle_id}/history")
async def vehicle_history(vehicle_id: str, user: dict = Depends(get_current_user)):
    entries = await db.audit.find({"vehicle_id": vehicle_id}).sort("ts", -1).to_list(1000)
    for e in entries:
        e.pop("_id", None)
        e.pop("restore", None)
    return entries

@api_router.post("/audit/{audit_id}/undo")
async def undo_operation(audit_id: str, user: dict = Depends(require("delete_operations"))):
    entry = await db.audit.find_one({"id": audit_id})
    if not entry:
        raise HTTPException(status_code=404, detail="Operazione non trovata")
    if entry.get("reverted"):
        raise HTTPException(status_code=400, detail="Operazione già annullata")
    restore = entry.get("restore")
    if not restore:
        raise HTTPException(status_code=400, detail="Operazione non annullabile")
    vid = entry.get("vehicle_id")
    rtype = restore.get("type")
    if rtype == "set":
        if not await db.vehicles.find_one({"id": vid}):
            raise HTTPException(status_code=400, detail="Veicolo non più presente")
        await db.vehicles.update_one({"id": vid}, {"$set": restore["set"]})
    elif rtype == "delete_vehicle":
        await db.vehicles.delete_one({"id": vid})
    elif rtype == "insert_vehicle":
        doc = dict(restore["doc"])
        doc.pop("_id", None)
        if not await db.vehicles.find_one({"id": vid}):
            await db.vehicles.insert_one(doc)
    await db.audit.update_one({"id": audit_id}, {"$set": {
        "reverted": True, "reverted_by": user["email"], "reverted_at": now_iso()}})
    await log_op(user, "undo", "Annullamento operazione",
                 {"id": vid, "targa": entry.get("targa"), "marca_modello": entry.get("marca_modello")},
                 f"Annullata operazione: {entry.get('action_label')} — {entry.get('description')}", None)
    return {"ok": True}

# ---------------- Dashboard & calendar ----------------
@api_router.get("/login-log")
async def login_log(limit: int = Query(200), user: dict = Depends(require("manage_users"))):
    logs = await db.login_log.find().sort("ts", -1).to_list(min(limit, 1000))
    for l in logs:
        l.pop("_id", None)
    return logs

@api_router.get("/notifications/today")
async def notifications_today(user: dict = Depends(get_current_user)):
    settings = await get_settings()
    bell_days = settings.get("bell_days", 7)
    vehicles = await db.vehicles.find().to_list(2000)
    today = datetime.now(timezone.utc).date()
    overdue, due_today, upcoming = [], [], []
    for v in vehicles:
        if v.get("policy"):
            normalize_policy_suspension(v["policy"])
        for ev in vehicle_events(build_vehicle_view(v)):
            d = parse_date(ev["date"])
            if not d:
                continue
            dl = (d - today).days
            ev["days_left"] = dl
            if dl < 0:
                overdue.append(ev)
            elif dl == 0:
                due_today.append(ev)
            elif dl <= bell_days:
                upcoming.append(ev)
    for arr in (overdue, due_today, upcoming):
        arr.sort(key=lambda e: e["date"])
    return {"overdue": overdue, "today": due_today, "upcoming": upcoming,
            "bell_days": bell_days, "count": len(overdue) + len(due_today)}

async def fetch_audit(vehicle_id, action, date_from, date_to, user_email=None, limit=5000):
    q = build_audit_query(vehicle_id, action, date_from, date_to, user_email)
    entries = await db.audit.find(q).sort("ts", -1).to_list(limit)
    for e in entries:
        e.pop("_id", None)
        e.pop("restore", None)
    return entries

@api_router.get("/audit/operators")
async def audit_operators(user: dict = Depends(get_current_user)):
    emails = await db.audit.distinct("user_email")
    out = []
    for em in emails:
        if not em:
            continue
        one = await db.audit.find_one({"user_email": em}, sort=[("ts", -1)])
        out.append({"email": em, "name": (one or {}).get("user_name") or em})
    return out

@api_router.get("/reports/audit/excel")
async def audit_report_excel(vehicle_id: Optional[str] = Query(None), action: Optional[str] = Query(None),
                             user_email: Optional[str] = Query(None),
                             date_from: Optional[str] = Query(None), date_to: Optional[str] = Query(None),
                             user: dict = Depends(require("export_reports"))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    settings = await get_settings()
    entries = await fetch_audit(vehicle_id, action, date_from, date_to, user_email)
    wb = Workbook(); ws = wb.active; ws.title = "Storico"
    ws["A1"] = settings.get("company_name", "FleetCare Autonoleggio"); ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Resoconto storico operazioni"
    headers = ["Data", "Ora", "Operazione", "Targa", "Veicolo", "Descrizione", "Operatore", "Stato"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=ci, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    for ri, e in enumerate(entries, 5):
        ts = e.get("ts", "")
        vals = [e.get("effective_date", "")[:10], ts[11:16] if len(ts) > 16 else "",
                e.get("action_label", ""), e.get("targa") or "-", e.get("marca_modello") or "-",
                e.get("description", ""), e.get("user_name") or e.get("user_email", ""),
                "ANNULLATA" if e.get("reverted") else ""]
        for ci, val in enumerate(vals, 1):
            ws.cell(row=ri, column=ci, value=val)
    widths = [12, 8, 24, 12, 20, 50, 22, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"storico_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={fname}"})

@api_router.get("/reports/audit/pdf")
async def audit_report_pdf(vehicle_id: Optional[str] = Query(None), action: Optional[str] = Query(None),
                           user_email: Optional[str] = Query(None),
                           date_from: Optional[str] = Query(None), date_to: Optional[str] = Query(None),
                           user: dict = Depends(require("export_reports"))):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    settings = await get_settings()
    entries = await fetch_audit(vehicle_id, action, date_from, date_to, user_email)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=1*cm, bottomMargin=1*cm, leftMargin=1*cm, rightMargin=1*cm)
    styles = getSampleStyleSheet()
    period = ""
    if date_from or date_to:
        period = f" · periodo {date_from or '…'} -> {date_to or '…'}"
    elements = [Paragraph(settings.get("company_name", "FleetCare Autonoleggio"), styles["Title"]),
                Paragraph("Resoconto storico operazioni" + period, styles["Heading2"]),
                Paragraph(f"Generato il {datetime.now().strftime('%d/%m/%Y %H:%M')} · {len(entries)} operazioni", styles["Normal"]),
                Spacer(1, 0.4*cm)]
    desc_style = styles["Normal"]; desc_style.fontSize = 7
    cols = ["Data", "Operazione", "Targa", "Descrizione", "Operatore", "Stato"]
    data = [cols]
    for e in entries:
        data.append([e.get("effective_date", "")[:10], e.get("action_label", ""), e.get("targa") or "-",
                     Paragraph(escape(e.get("description", "")), desc_style),
                     e.get("user_name") or e.get("user_email", ""),
                     "ANNULLATA" if e.get("reverted") else ""])
    table = Table(data, repeatRows=1, colWidths=[2.2*cm, 3.5*cm, 2*cm, 11*cm, 4*cm, 2.3*cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")])]))
    elements.append(table)
    doc.build(elements)
    buf.seek(0)
    fname = f"storico_{datetime.now().strftime('%Y%m%d')}.pdf"
    return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={fname}"})

@api_router.get("/reports/vehicle/{vehicle_id}/pdf")
async def vehicle_report_pdf(vehicle_id: str, user: dict = Depends(require("export_reports"))):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
    from reportlab.lib.styles import getSampleStyleSheet
    v = await get_vehicle_or_404(vehicle_id)
    view = build_vehicle_view(v)
    settings = await get_settings()
    entries = await db.audit.find({"vehicle_id": vehicle_id}).sort("ts", -1).to_list(1000)
    for e in entries:
        e.pop("_id", None)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.2*cm, bottomMargin=1.2*cm, leftMargin=1.5*cm, rightMargin=1.5*cm)
    styles = getSampleStyleSheet()
    small = styles["Normal"]; small.fontSize = 8
    elements = []
    if settings.get("logo_path"):
        try:
            content, _ = get_object(settings["logo_path"])
            if is_valid_image_bytes(content):
                elements.append(Image(io.BytesIO(content), width=1.5*cm, height=1.5*cm))
        except Exception:
            pass
    elements += [Paragraph(settings.get("company_name", "FleetCare Autonoleggio"), styles["Title"]),
                 Paragraph(f"Scheda veicolo — {view['targa']}", styles["Heading2"]),
                 Paragraph(f"Generata il {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"]),
                 Spacer(1, 0.3*cm)]

    def dt(s):
        return s[:10] if s else "—"

    can = "PUO CIRCOLARE" if view["can_circulate"] else "NON PUO CIRCOLARE"
    gen = [["Targa", view["targa"], "Marca/Modello", view.get("marca_modello", "")],
           ["Immatricolazione", dt(view.get("data_immatricolazione")), "Circolazione", can],
           ["Scad. bollo", dt(view.get("bollo_scadenza")), "Scad. collaudo", dt(view.get("collaudo_deadline"))]]
    p = view.get("policy")
    if p:
        gen += [["Compagnia", p.get("compagnia") or "—", "N. polizza", p.get("numero_polizza") or "—"],
                ["Tipo polizza", p.get("tipologia") or "—", "Scad. contratto", dt(p.get("scadenza_contratto"))],
                ["Premio", (f"€ {p['importo_premio']:.2f}" if p.get("importo_premio") else "—"),
                 "Sospensioni", f"{p.get('cumulative_suspension_days', 0)}/{MAX_SUSPENSION_DAYS} gg"],
                ["Comporto 15gg", ("Si" if p.get("grace_applies") else "No"), "Stato polizza", p.get("status") or "—"]]
    t = Table(gen, colWidths=[3.5*cm, 5.5*cm, 3.5*cm, 5*cm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                           ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1F5F9")),
                           ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F1F5F9")),
                           ("FONTSIZE", (0, 0), (-1, -1), 8), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    elements += [Paragraph("Dati generali", styles["Heading3"]), t, Spacer(1, 0.3*cm)]
    if view.get("note"):
        elements += [Paragraph("Note", styles["Heading3"]),
                     Paragraph(escape(view["note"]), small), Spacer(1, 0.3*cm)]

    def hist_table(title, header, rows):
        if not rows:
            return [Paragraph(title, styles["Heading3"]), Paragraph("Nessun dato.", small), Spacer(1, 0.2*cm)]
        data = [header] + rows
        tb = Table(data, repeatRows=1)
        tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 7),
                                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
                                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")])]))
        return [Paragraph(title, styles["Heading3"]), tb, Spacer(1, 0.3*cm)]

    bollo_rows = [[dt(b.get("data_pagamento")), f"€ {b.get('importo', 0):.2f}", b.get("periodo") or "—", b.get("note") or "—"]
                  for b in view.get("bollo_history", [])]
    elements += hist_table("Storico bolli", ["Data", "Importo", "Periodo", "Note"], bollo_rows)

    coll_rows = [[dt(c.get("data_collaudo")), c.get("note") or "—"] for c in view.get("collaudo_history", [])]
    elements += hist_table("Storico collaudi", ["Data", "Note"], coll_rows)

    op_rows = [[dt(e.get("effective_date")), e.get("action_label", ""),
                Paragraph(escape(e.get("description", "")), small), e.get("user_name") or e.get("user_email", ""),
                "ANNULLATA" if e.get("reverted") else ""] for e in entries]
    if op_rows:
        data = [["Data", "Operazione", "Descrizione", "Operatore", "Stato"]] + op_rows
        tb = Table(data, repeatRows=1, colWidths=[2.2*cm, 3*cm, 7.3*cm, 3.5*cm, 2*cm])
        tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 7),
                                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")])]))
        elements += [Paragraph("Scheda storica operazioni", styles["Heading3"]), tb]
    else:
        elements += [Paragraph("Scheda storica operazioni", styles["Heading3"]), Paragraph("Nessuna operazione.", small)]

    doc.build(elements)
    buf.seek(0)
    fname = f"scheda_{view['targa']}_{datetime.now().strftime('%Y%m%d')}.pdf"
    return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={fname}"})

@api_router.get("/dashboard/stats")
async def dashboard_stats(user: dict = Depends(get_current_user)):
    vehicles = await db.vehicles.find().to_list(2000)
    total = can = cannot = suspended = upcoming = 0
    for v in vehicles:
        if v.get("policy"):
            normalize_policy_suspension(v["policy"])
        view = build_vehicle_view(v)
        total += 1
        if view["can_circulate"]:
            can += 1
        else:
            cannot += 1
        if view["policy"] and view["policy"]["status"] == "suspended":
            suspended += 1
        states = [view["collaudo_state"], view["bollo_state"]]
        if view["policy"]:
            states += [view["insurance_state"], view["policy"]["rata_state"]]
        if any(s == "upcoming" for s in states):
            upcoming += 1
    return {"total": total, "can_circulate": can, "cannot_circulate": cannot,
            "suspended": suspended, "upcoming_30": upcoming}

# ---------------- Strategia sospensioni ----------------
PEAK_MONTHS = {7, 8}
PHASE_LABEL = {"closed": "Bassa stagione (attività chiusa)",
               "shoulder": "Stagione di spalla", "peak": "Alta stagione (picco luglio–agosto)"}
PHASE_DEMAND = {"closed": 90.0, "shoulder": 45.0, "peak": 0.0}
TIPO_WEIGHT = {"auto": 0.7, "furgone": 1.0, "altro": 1.0}

def season_phase(d: date, start_m: int, end_m: int) -> str:
    m = d.month
    in_season = (start_m <= m <= end_m) if start_m <= end_m else (m >= start_m or m <= end_m)
    if not in_season:
        return "closed"
    return "peak" if m in PEAK_MONTHS else "shoulder"

def next_peak_prep(today: date) -> date:
    cand = date(today.year, 7, 1) - timedelta(days=15)
    return cand if today < cand else date(today.year + 1, 7, 1) - timedelta(days=15)

def strategy_for_vehicle(view: dict, phase: str, today: date, start_m: int) -> dict:
    p = view.get("policy")
    tipo = view.get("tipo") or "auto"
    base = {"vehicle_id": view["id"], "targa": view["targa"], "marca_modello": view.get("marca_modello"),
            "tipo": tipo, "can_circulate": view.get("can_circulate")}
    if not p:
        return {**base, "score": 0, "level": "none", "recommendation": "Nessuna polizza",
                "reasons": ["Nessuna polizza assicurativa: non c'è copertura da sospendere."],
                "blockers": [], "estimated_saving": None, "suggested_from": None,
                "suggested_until": None, "policy": None}

    contract = parse_date(p.get("scadenza_contratto"))
    blockers = []
    suspended = p.get("status") == "suspended"
    if suspended:
        blockers.append("Polizza già sospesa")
    elif not p.get("can_suspend"):
        blockers.append("Limite di 10 mesi di sospensione raggiunto" if p.get("suspension_limit_reached")
                        else "Sospensione non disponibile")
    in_grace_block = False
    if contract and contract < today <= contract + timedelta(days=GRACE_DAYS):
        in_grace_block = True
        blockers.append("Nei 15 giorni di comporto di legge: sospensione non consentita")

    cost_events = []
    if view.get("collaudo_deadline"):
        cost_events.append(("Collaudo/revisione", parse_date(view["collaudo_deadline"]), 120.0))
    if contract:
        cost_events.append(("Rinnovo polizza", contract, p.get("importo_premio") or 300.0))
    rata = parse_date(p.get("scadenza_rata_intermedia"))
    if rata:
        cost_events.append(("Rata polizza", rata, p.get("importo_rata") or 150.0))
    upcoming = [(lbl, dt, amt, (dt - today).days) for (lbl, dt, amt) in cost_events
                if dt and -15 <= (dt - today).days <= 210]
    upcoming.sort(key=lambda x: x[3])

    reasons = []
    if phase == "closed":
        reasons.append("Siamo in bassa stagione (attività chiusa): la richiesta di veicoli è minima.")
    elif phase == "peak":
        reasons.append("Alta stagione (luglio–agosto): il veicolo è quasi certamente necessario.")
    else:
        reasons.append("Stagione di spalla: richiesta moderata, valuta caso per caso.")

    if tipo != "auto" and phase != "peak":
        reasons.append(f"Fuori dal picco servono soprattutto auto: questo mezzo ({tipo}) è un buon candidato alla sospensione.")
    elif tipo == "auto":
        reasons.append("È un'auto: potrebbe servire anche fuori dal picco, valuta la reale necessità.")

    econ = 0.0
    if upcoming:
        lbl, dt, amt, dd = upcoming[0]
        proximity = max(0.0, 100.0 * (1 - max(dd, 0) / 180.0))
        amount_factor = min(1.0, (amt or 0) / 1000.0)
        econ = proximity * (0.6 + 0.4 * amount_factor)
        if dd < 0:
            reasons.append(f"{lbl} già scaduto da {abs(dd)} gg (~€{amt:.0f}).")
        else:
            reasons.append(f"{lbl} in scadenza tra {dd} gg (~€{amt:.0f}): sospendendo ora rimandi l'esborso a stagione avviata.")
        if dt.month in (start_m, (start_m - 1) or 12):
            reasons.append("L'esborso cade a inizio stagione, periodo di scarsa liquidità: meglio pianificarlo ad attività avviata.")

    type_demand = min(100.0, PHASE_DEMAND.get(phase, 45.0) * TIPO_WEIGHT.get(tipo, 1.0))
    score = 0.55 * econ + 0.45 * type_demand
    if phase == "peak":
        score *= 0.2
    if blockers:
        score = min(score, 15.0)
    score = int(round(min(100.0, max(0.0, score))))

    sug_from = today if not in_grace_block else (contract + timedelta(days=GRACE_DAYS + 1))
    sug_until = next_peak_prep(today)
    remaining = MAX_SUSPENSION_DAYS - p.get("cumulative_suspension_days", 0)
    if sug_until <= sug_from:
        sug_until = sug_from + timedelta(days=min(90, max(remaining, 0)))
    if (sug_until - sug_from).days > remaining:
        sug_until = sug_from + timedelta(days=max(remaining, 0))

    estimated_saving = None
    if p.get("importo_premio") and sug_until > sug_from:
        estimated_saving = round(p["importo_premio"] * (sug_until - sug_from).days / 365.0, 2)

    if suspended:
        level, recommendation = "suspended", "Già sospesa"
    elif not p.get("can_suspend"):
        level, recommendation = "low", "Non sospendibile"
    elif score >= 60:
        level, recommendation = "high", "Sospendi ora"
    elif score >= 35:
        level, recommendation = "medium", "Valuta sospensione"
    else:
        level, recommendation = "low", "Mantieni attivo"

    return {**base, "score": score, "level": level, "recommendation": recommendation,
            "reasons": reasons, "blockers": blockers, "estimated_saving": estimated_saving,
            "suggested_from": sug_from.isoformat(), "suggested_until": sug_until.isoformat(),
            "policy": {"compagnia": p.get("compagnia"), "scadenza_contratto": p.get("scadenza_contratto"),
                       "status": p.get("status"),
                       "cumulative_suspension_days": p.get("cumulative_suspension_days", 0),
                       "max_suspension_days": p.get("max_suspension_days", MAX_SUSPENSION_DAYS),
                       "can_suspend": p.get("can_suspend", False),
                       "importo_premio": p.get("importo_premio")}}

@api_router.get("/strategy")
async def strategy(user: dict = Depends(get_current_user)):
    settings = await get_settings()
    start_m = int(settings.get("season_start_month", 4))
    end_m = int(settings.get("season_end_month", 10))
    today = datetime.now(timezone.utc).date()
    phase = season_phase(today, start_m, end_m)
    vehicles = await db.vehicles.find().to_list(2000)
    items = []
    for v in vehicles:
        if v.get("policy"):
            normalize_policy_suspension(v["policy"])
        items.append(strategy_for_vehicle(build_vehicle_view(v), phase, today, start_m))
    items.sort(key=lambda x: x["score"], reverse=True)
    opportunities = sum(1 for i in items if i["level"] == "high")
    return {"today": today.isoformat(),
            "season": {"start": start_m, "end": end_m, "phase": phase, "phase_label": PHASE_LABEL.get(phase)},
            "items": items, "opportunities": opportunities}

def vehicle_events(view: dict) -> List[dict]:
    ev = []
    if view.get("bollo_scadenza"):
        ev.append({"date": view["bollo_scadenza"], "type": "bollo", "label": "Bollo",
                   "state": view["bollo_state"], "targa": view["targa"], "vehicle_id": view["id"],
                   "marca_modello": view["marca_modello"]})
    if view.get("collaudo_deadline"):
        ev.append({"date": view["collaudo_deadline"], "type": "collaudo", "label": "Collaudo",
                   "state": view["collaudo_state"], "targa": view["targa"], "vehicle_id": view["id"],
                   "marca_modello": view["marca_modello"]})
    p = view.get("policy")
    if p:
        if p.get("scadenza_rata_intermedia"):
            ev.append({"date": p["scadenza_rata_intermedia"], "type": "rata", "label": "Rata polizza",
                       "state": p["rata_state"], "targa": view["targa"], "vehicle_id": view["id"],
                       "marca_modello": view["marca_modello"]})
        if p.get("scadenza_contratto"):
            ev.append({"date": p["scadenza_contratto"], "type": "polizza", "label": "Contratto polizza",
                       "state": view["insurance_state"], "targa": view["targa"], "vehicle_id": view["id"],
                       "marca_modello": view["marca_modello"]})
    return ev

@api_router.get("/deadlines")
async def deadlines(user: dict = Depends(get_current_user)):
    vehicles = await db.vehicles.find().to_list(2000)
    events = []
    for v in vehicles:
        if v.get("policy"):
            normalize_policy_suspension(v["policy"])
        events += vehicle_events(build_vehicle_view(v))
    return events

# ---------------- Reports ----------------
def gather_report_rows(vehicles):
    rows = []
    for v in vehicles:
        if v.get("policy"):
            normalize_policy_suspension(v["policy"])
        view = build_vehicle_view(v)
        p = view["policy"]
        rows.append({
            "Targa": view.get("targa", ""), "Marca/Modello": view.get("marca_modello", ""),
            "Immatricolazione": view.get("data_immatricolazione", ""),
            "Circolazione": "PUO CIRCOLARE" if view["can_circulate"] else "NON PUO CIRCOLARE",
            "Scad. Bollo": view.get("bollo_scadenza") or "-",
            "Scad. Collaudo": view.get("collaudo_deadline") or "-",
            "Compagnia": (p["compagnia"] if p else "-"),
            "N. Polizza": (p.get("numero_polizza") or "-" if p else "-"),
            "Tipo Polizza": (p["tipologia"] if p else "-"),
            "Premio": (f'{p["importo_premio"]:.2f}' if p and p.get("importo_premio") else "-"),
            "Scad. Contratto": (p["scadenza_contratto"] if p else "-"),
            "Comporto 15gg": ("Si" if p and p["grace_applies"] else "No") if p else "-",
            "Stato Polizza": (p["status"] if p else "-"),
            "GG Sospensione": (f'{p["cumulative_suspension_days"]}/{MAX_SUSPENSION_DAYS}' if p else "-"),
        })
    return rows

@api_router.get("/reports/excel")
async def report_excel(user: dict = Depends(require("export_reports"))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.drawing.image import Image as XLImage
    settings = await get_settings()
    vehicles = await db.vehicles.find().sort("targa", 1).to_list(2000)
    rows = gather_report_rows(vehicles)
    wb = Workbook(); ws = wb.active; ws.title = "Flotta"
    if settings.get("logo_path"):
        try:
            content, _ = get_object(settings["logo_path"])
            if not is_valid_image_bytes(content):
                raise ValueError("logo non valido")
            img = XLImage(io.BytesIO(content)); img.height = 48; img.width = 48
            ws.add_image(img, "A1")
            ws["B1"] = settings.get("company_name", "FleetCare"); ws["B1"].font = Font(bold=True, size=14)
            ws.row_dimensions[1].height = 40
        except Exception as e:
            ws["A1"] = settings.get("company_name", "FleetCare Autonoleggio"); ws["A1"].font = Font(bold=True, size=14)
    else:
        ws["A1"] = settings.get("company_name", "FleetCare Autonoleggio"); ws["A1"].font = Font(bold=True, size=14)
    headers = list(rows[0].keys()) if rows else ["Targa", "Marca/Modello"]
    hr = 3
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=hr, column=ci, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    for ri, r in enumerate(rows, hr + 1):
        for ci, h in enumerate(headers, 1):
            ws.cell(row=ri, column=ci, value=r[h])
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[chr(64 + i)].width = max(14, len(h) + 4)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"flotta_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={fname}"})

@api_router.get("/reports/pdf")
async def report_pdf(user: dict = Depends(require("export_reports"))):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
    from reportlab.lib.styles import getSampleStyleSheet
    settings = await get_settings()
    vehicles = await db.vehicles.find().sort("targa", 1).to_list(2000)
    rows = gather_report_rows(vehicles)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=1*cm, bottomMargin=1*cm, leftMargin=1*cm, rightMargin=1*cm)
    styles = getSampleStyleSheet()
    elements = []
    if settings.get("logo_path"):
        try:
            content, _ = get_object(settings["logo_path"])
            if is_valid_image_bytes(content):
                elements.append(Image(io.BytesIO(content), width=1.6*cm, height=1.6*cm))
        except Exception as e:
            logger.warning(f"logo pdf skip: {e}")
    elements += [Paragraph(settings.get("company_name", "FleetCare Autonoleggio"), styles["Title"]),
                 Paragraph("Scadenziario Flotta Veicoli", styles["Heading2"]),
                 Paragraph(f"Generato il {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"]), Spacer(1, 0.4*cm)]
    cols = ["Targa", "Marca/Modello", "Circolazione", "Scad. Bollo", "Scad. Collaudo",
            "Compagnia", "N. Polizza", "Premio", "Scad. Contratto", "Comporto 15gg", "Stato Polizza"]
    data = [cols] + [[str(r[c]) for c in cols] for r in rows]
    table = Table(data, repeatRows=1)
    style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
             ("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")])]
    for i, r in enumerate(rows, 1):
        col = colors.HexColor("#E11D48") if r["Circolazione"] == "NON PUO CIRCOLARE" else colors.HexColor("#059669")
        style.append(("TEXTCOLOR", (2, i), (2, i), col))
    table.setStyle(TableStyle(style))
    elements.append(table)
    doc.build(elements)
    buf.seek(0)
    fname = f"scadenziario_{datetime.now().strftime('%Y%m%d')}.pdf"
    return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={fname}"})

# ---------------- Settings ----------------
@api_router.get("/settings")
async def read_settings(user: dict = Depends(get_current_user)):
    return await get_settings()

@api_router.put("/settings")
async def update_settings(input: SettingsInput, user: dict = Depends(require("manage_settings"))):
    recipients = [e.strip().lower() for e in input.notification_recipients if e.strip()]
    nd = input.notification_days
    await db.settings.update_one({"key": "app"}, {"$set": {
        "company_name": input.company_name or "FleetCare Autonoleggio",
        "notification_recipients": recipients,
        "notification_days": {"bollo": max(1, nd.bollo), "collaudo": max(1, nd.collaudo), "polizza": max(1, nd.polizza)},
        "bell_days": max(0, int(input.bell_days)),
        "season_start_month": int(input.season_start_month),
        "season_end_month": int(input.season_end_month),
    }}, upsert=True)
    return await get_settings()

@api_router.post("/settings/logo")
async def upload_logo(file: UploadFile = File(...), user: dict = Depends(require("manage_settings"))):
    ext = (file.filename.rsplit(".", 1)[-1] if "." in file.filename else "png").lower()
    if ext not in {"png", "jpg", "jpeg", "webp"}:
        raise HTTPException(status_code=400, detail="Il logo deve essere PNG, JPG o WEBP")
    data = await file.read()
    if not is_valid_image_bytes(data):
        raise HTTPException(status_code=400, detail="File immagine non valido o corrotto")
    path = f"{APP_NAME}/branding/logo.{ext}"
    result = put_object(path, data, MIME_TYPES.get(ext, "image/png"))
    await db.settings.update_one({"key": "app"}, {"$set": {"logo_path": result["path"]}}, upsert=True)
    return {"logo_path": result["path"]}

# ---------------- Notifications ----------------
async def compute_digest():
    settings = await get_settings()
    nd = settings.get("notification_days", {"bollo": 30, "collaudo": 30, "polizza": 30})
    type_key = {"bollo": "bollo", "collaudo": "collaudo", "polizza": "polizza", "rata": "polizza"}
    vehicles = await db.vehicles.find().to_list(2000)
    today = datetime.now(timezone.utc).date()
    items = []
    for v in vehicles:
        if v.get("policy"):
            normalize_policy_suspension(v["policy"])
        for ev in vehicle_events(build_vehicle_view(v)):
            d = parse_date(ev["date"])
            if not d:
                continue
            window = nd.get(type_key.get(ev["type"], "polizza"), 30)
            days_left = (d - today).days
            if days_left <= window:
                ev["days_left"] = days_left
                items.append(ev)
    items.sort(key=lambda e: e["date"])
    return items, settings

def digest_html(company: str, items: List[dict]) -> str:
    rows = ""
    for it in items:
        overdue = it["days_left"] < 0
        color = "#E11D48" if overdue else ("#B45309" if it["days_left"] <= 7 else "#0F172A")
        gg = f'{abs(it["days_left"])} gg fa' if overdue else f'tra {it["days_left"]} gg'
        rows += (f'<tr><td style="padding:6px 10px;border-bottom:1px solid #eee">{escape(it["targa"])}</td>'
                 f'<td style="padding:6px 10px;border-bottom:1px solid #eee">{escape(it["marca_modello"])}</td>'
                 f'<td style="padding:6px 10px;border-bottom:1px solid #eee">{escape(it["label"])}</td>'
                 f'<td style="padding:6px 10px;border-bottom:1px solid #eee">{escape(it["date"])}</td>'
                 f'<td style="padding:6px 10px;border-bottom:1px solid #eee;color:{color};font-weight:bold">{gg}</td></tr>')
    if not rows:
        rows = '<tr><td colspan="5" style="padding:12px;color:#888">Nessuna scadenza in arrivo.</td></tr>'
    return (f'<table role="presentation" width="100%"><tr><td style="padding:20px;font-family:Arial,sans-serif">'
            f'<h2 style="color:#0F172A;margin:0 0 4px">{escape(company)}</h2>'
            f'<p style="color:#475569;margin:0 0 16px">Riepilogo scadenze flotta</p>'
            f'<table width="100%" style="border-collapse:collapse;font-size:13px">'
            f'<tr style="background:#0F172A;color:#fff"><th style="padding:6px 10px;text-align:left">Targa</th>'
            f'<th style="padding:6px 10px;text-align:left">Veicolo</th><th style="padding:6px 10px;text-align:left">Scadenza</th>'
            f'<th style="padding:6px 10px;text-align:left">Data</th><th style="padding:6px 10px;text-align:left">Quando</th></tr>{rows}</table>'
            f'<p style="font-size:12px;color:#888;margin-top:16px">Inviato da {escape(EMAIL_FROM_NAME)}. '
            f'Non chiediamo mai password o dati di pagamento via email.</p></td></tr></table>')

_last_digest_at = 0.0

async def run_digest():
    global _last_digest_at
    import time as _time
    now = _time.monotonic()
    if now - _last_digest_at < 15:
        items, _ = await compute_digest()
        return {"sent": 0, "reason": "riepilogo inviato da poco, riprova tra qualche secondo", "items": len(items)}
    _last_digest_at = now
    items, settings = await compute_digest()
    recipients = settings.get("notification_recipients", [])
    if not recipients:
        return {"sent": 0, "reason": "nessun destinatario configurato", "items": len(items)}
    company = settings.get("company_name", "FleetCare Autonoleggio")
    html = digest_html(company, items)
    subject = f"Scadenze flotta — {len(items)} in arrivo ({datetime.now().strftime('%d/%m/%Y')})"
    sent = 0
    for r in recipients:
        try:
            await send_email(to=r, subject=subject, html=html)
            sent += 1
        except Exception as e:
            logger.error(f"digest send to {r} failed: {e}")
    return {"sent": sent, "recipients": recipients, "items": len(items)}

@api_router.post("/notifications/send-now")
async def send_now(user: dict = Depends(require("manage_settings"))):
    return await run_digest()

@api_router.get("/notifications/preview")
async def notifications_preview(user: dict = Depends(get_current_user)):
    items, settings = await compute_digest()
    return {"count": len(items), "items": items, "recipients": settings.get("notification_recipients", [])}

@api_router.post("/cron/deadlines-digest")
async def cron_digest(background_tasks: BackgroundTasks, authorization: str = Header(None)):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    secret = os.environ.get("WEBHOOK_CRON_SECRET", "")
    provided = authorization[7:] if authorization and authorization.startswith("Bearer ") else ""
    if not secret or not pysecrets.compare_digest(provided, secret):
        raise HTTPException(status_code=401, detail="Unauthorized")
    background_tasks.add_task(run_digest)
    return {"accepted": True}

@api_router.get("/")
async def root():
    return {"message": "FleetCare API"}

app.include_router(api_router)
app.add_middleware(CORSMiddleware, allow_credentials=True,
                   allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
                   allow_methods=["*"], allow_headers=["*"])

@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.vehicles.create_index("id", unique=True)
    await db.audit.create_index("ts")
    await db.audit.create_index("vehicle_id")
    await db.login_log.create_index("ts")
    try:
        init_storage()
    except Exception as e:
        logger.error(f"Storage init failed: {e}")
    admin_email = os.environ.get("ADMIN_EMAIL", "").lower()
    admin_password = os.environ.get("ADMIN_PASSWORD", "")
    if admin_email and admin_password:
        existing = await db.users.find_one({"email": admin_email})
        if existing is None:
            await db.users.insert_one({"email": admin_email, "password_hash": hash_password(admin_password),
                                       "name": "Titolare", "role": "admin", "permissions": list(PERMISSIONS),
                                       "is_active": True, "created_at": now_iso()})
        else:
            upd = {"role": "admin", "is_active": True}
            if not verify_password(admin_password, existing["password_hash"]):
                upd["password_hash"] = hash_password(admin_password)
            await db.users.update_one({"email": admin_email}, {"$set": upd})
    await get_settings()
    await db.vehicles.update_many(
        {"policy.tipologia": "semestrale"},
        {"$set": {"policy.tipologia": "annuale", "policy.frazionamento": "semestrale"}})

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
